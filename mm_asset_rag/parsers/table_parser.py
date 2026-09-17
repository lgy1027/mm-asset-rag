"""Semantic row chunks for CSV, TSV, and Excel workbooks."""

from __future__ import annotations

import csv
from collections.abc import Iterator
from pathlib import Path

from ..core.schema import ParsedChunk
from ..core.settings import get_settings
from ..ingest.assets import IngestAsset


def parse_table(asset: IngestAsset) -> list[ParsedChunk]:
    """Turn each non-empty table row into header-qualified retrieval text."""
    settings = get_settings()
    chunks: list[ParsedChunk] = []
    for sheet_name, rows in _read_sheets(asset.file_path):
        row_iterator = iter(rows)
        header = next(row_iterator, None)
        if header is None:
            continue
        if len(header) > settings.table_max_columns:
            raise ValueError("table column budget exceeded")
        headers = [value.strip() or f"列{index}" for index, value in enumerate(header, start=1)]
        total_chars = 0
        for row_index, row in enumerate(row_iterator, start=1):
            if row_index > settings.table_max_rows:
                raise ValueError("table row budget exceeded")
            if len(row) > len(headers):
                raise ValueError("table row has more values than headers")
            values = [value.strip() for value in row]
            if any(len(value) > settings.table_max_cell_chars for value in values):
                raise ValueError("table cell budget exceeded")
            if not any(values):
                continue
            pairs = [
                f"{header_name}：{value}" for header_name, value in zip(headers, values) if value
            ]
            if not pairs:
                continue
            prefix = f"表格：{Path(asset.relative_path).stem}"
            if sheet_name:
                prefix += f"；工作表：{sheet_name}"
            text = "\n".join([prefix, *pairs])
            total_chars += len(text)
            if total_chars > settings.table_max_total_chars:
                raise ValueError("table text budget exceeded")
            chunks.append(
                ParsedChunk(
                    text=text,
                    metadata={
                        "asset_id": asset.asset_id,
                        "asset_title": asset.title,
                        "source_type": asset.source_type,
                        "source_path": asset.relative_path,
                        "source_url": asset.source_url,
                        "parser": "table-semantic",
                        "table_headers": headers,
                        "sheet_name": sheet_name,
                        "row_index": row_index,
                    },
                )
            )
    return chunks


def _read_sheets(path: Path) -> Iterator[tuple[str, Iterator[list[str]]]]:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        with path.open(encoding="utf-8-sig", newline="") as file_obj:
            rows = (
                [str(value) for value in row] for row in csv.reader(file_obj, delimiter=delimiter)
            )
            yield "", rows
        return
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                rows = (
                    ["" if value is None else str(value) for value in row]
                    for row in worksheet.iter_rows(values_only=True)
                )
                yield worksheet.title, rows
        finally:
            workbook.close()
        return
    raise ValueError(f"unsupported table format: {path.suffix}")
